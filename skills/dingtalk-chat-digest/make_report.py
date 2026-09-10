# -*- coding: utf-8 -*-
"""
钉钉聊天汇总报告生成器 (dingtalk-chat-digest)
把「逐条原文 + AI总结」排版成 Word 报告。

用法:
  python make_report.py \
    --title "我与<姓名>近三天单聊记录 · 原文与总结" \
    --meta  "时间范围：2026-08-18 至 2026-08-20 ｜ 会话：单聊 ｜ 消息数：5 条" \
    --messages messages.json \
    --summary  summary.json \
    --out "<输出目录>/<姓名>聊天汇总.docx" \
    [--self "<本人姓名或'我'>"]   # 我方姓名,用于原文着色区分(可选)

messages.json: dws +chat-messages 返回里的 messages 数组
  每项至少含 time(或 createTime)、sender、text
summary.json: {"sections":[{"heading":..,"points":[..]}], "oneLine":.., "note":..}
"""
import argparse, json, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

NAVY = (0x1F, 0x3A, 0x5F)
ORANGE = (0xC0, 0x50, 0x00)
GRAY = (0x88, 0x88, 0x88)
LGRAY = (0x99, 0x99, 0x99)


def add_run(p, text, size=11, bold=False, color=None, name='微软雅黑'):
    r = p.add_run(text)
    r.font.name = name
    r._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--title', required=True)
    ap.add_argument('--meta', default='')
    ap.add_argument('--messages', required=True)
    ap.add_argument('--summary', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--self', dest='selfname', default='我')
    args = ap.parse_args()

    with open(args.messages, 'r', encoding='utf-8') as f:
        raw = json.load(f)
    # 兼容:直接是数组,或含 messages 键的完整返回
    msgs = raw.get('messages', raw) if isinstance(raw, dict) else raw
    with open(args.summary, 'r', encoding='utf-8') as f:
        summary = json.load(f)

    # 规整字段 + 按时间正序
    def gettime(m):
        return m.get('time') or m.get('createTime') or ''
    norm = [{'time': gettime(m), 'sender': m.get('sender', ''),
             'text': m.get('text', '')} for m in msgs]
    norm.sort(key=lambda x: x['time'])

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.font.size = Pt(11)

    # 标题
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(t, args.title, size=16, bold=True, color=NAVY)
    if args.meta:
        sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(sub, args.meta, size=9, color=GRAY)
    doc.add_paragraph()

    # 一、原文
    hp = doc.add_paragraph()
    add_run(hp, '一、原文呈现（按时间正序）', size=13, bold=True, color=NAVY)
    for m in norm:
        p = doc.add_paragraph()
        add_run(p, f"[{m['time']}] ", size=9, color=LGRAY)
        is_self = (args.selfname and args.selfname in m['sender'])
        add_run(p, f"{m['sender']}：", size=11, bold=True,
                color=NAVY if is_self else ORANGE)
        add_run(p, m['text'], size=11)
    doc.add_paragraph()

    # 二、总结
    hp = doc.add_paragraph()
    add_run(hp, '二、内容总结', size=13, bold=True, color=NAVY)
    for sec in summary.get('sections', []):
        p = doc.add_paragraph()
        add_run(p, sec.get('heading', ''), size=11.5, bold=True, color=ORANGE)
        for pt in sec.get('points', []):
            bp = doc.add_paragraph(style='List Bullet')
            add_run(bp, pt, size=11)

    one = summary.get('oneLine')
    if one:
        doc.add_paragraph()
        hp = doc.add_paragraph()
        add_run(hp, '一句话概括', size=12, bold=True, color=NAVY)
        p = doc.add_paragraph(); add_run(p, one, size=11)

    note = summary.get('note')
    if note:
        doc.add_paragraph()
        hp = doc.add_paragraph()
        add_run(hp, '※ 边界说明', size=11, bold=True, color=(0xB0, 0x30, 0x30))
        p = doc.add_paragraph(); add_run(p, note, size=10.5, color=(0x66, 0x66, 0x66))

    doc.save(args.out)
    print(args.out)


if __name__ == '__main__':
    main()
